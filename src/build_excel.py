"""Build the Excel pack from validated pipeline outputs."""
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from src.config import ROOT,LABEL,ATTRITION_THRESHOLD
from src.quality_checks import check

def workbook(summary,dept,total,tenure,level,source,patterns,finding):
    wb=Workbook();ws=wb.active;ws.title='Summary'
    def source_sheet(name,df):
        w=wb.create_sheet(name);w.append([LABEL]);w.append(list(df.columns))
        for row in df.itertuples(index=False,name=None):w.append([None if pd.isna(v) else v for v in row])
        w.freeze_panes='B3';w.auto_filter.ref=f'A2:{get_column_letter(len(df.columns))}{len(df)+2}'
        return w
    cols=['report_month','Department','opening','closing','exits','average_headcount','attrition_rate','accepted','declined','offer_acceptance_rate','time_to_hire_days','avg_tenure_years','overtime_share','avg_monthly_income']
    dw=source_sheet('Department',dept[cols]);tw=source_sheet('Monthly',total.drop(columns=['Department'],errors='ignore'))
    source_sheet('Tenure',tenure);source_sheet('Income by level',level);source_sheet('Patterns',patterns)
    employee=source[['EmployeeNumber','Department','JobRole','JobLevel','MonthlyIncome','Attrition','OverTime','YearsAtCompany']]
    source_sheet('Employee lookup data',employee)
    source_sheet('Group output',summary)
    ws['A2']='Workforce & recruiting';ws['A3']=f'SIMULATED MONTHLY REPORT • {total.report_month.iloc[0]} to {total.report_month.iloc[-1]}'
    ws['A5']='Report month';ws['B5']=total.report_month.iloc[-1]
    ws['D5']='Attrition threshold';ws['E5']=ATTRITION_THRESHOLD;ws['E5'].number_format='0.0%'
    ws['D6']='Demo threshold, not an industry benchmark'
    wb.defined_names.add(DefinedName('ReportMonths',attr_text=f"'Monthly'!$A$3:$A${len(total)+2}"))
    dv=DataValidation(type='list',formula1='ReportMonths');ws.add_data_validation(dv);dv.add(ws['B5'])
    dv2=DataValidation(type='decimal',operator='between',formula1=0,formula2=1);ws.add_data_validation(dv2);dv2.add(ws['E5'])
    def lookup(col):
        c=get_column_letter(total.columns.get_loc(col)+1)
        expression=f'INDEX(Monthly!${c}$3:${c}${len(total)+2},MATCH($B$5,Monthly!$A$3:$A${len(total)+2},0))'
        return f'=IF(ISBLANK({expression}),"N/A",{expression})'
    for row,label,col,fmt in [(8,'Closing headcount','closing','#,##0'),(9,'Monthly attrition','attrition_rate','0.00%'),(10,'Offer acceptance','offer_acceptance_rate','0.00%'),(11,'Time to hire (days)','time_to_hire_days','0.0'),(12,'Overtime share','overtime_share','0.00%'),(13,'Average tenure (years)','avg_tenure_years','0.0')]:
        ws.cell(row,1,label);ws.cell(row,2,lookup(col)).number_format=fmt
    for j,v in enumerate(['Department','Opening','Closing','Exits','Attrition'],1):ws.cell(16,j,v)
    n=len(dept)+2
    for row,department in enumerate(sorted(dept.Department.unique()),17):
        ws.cell(row,1,department)
        for j,letter in [(2,'C'),(3,'D'),(4,'E')]:ws.cell(row,j,f'=SUMIFS(Department!${letter}$3:${letter}${n},Department!$A$3:$A${n},$B$5,Department!$B$3:$B${n},$A{row})')
        ws.cell(row,5,f'=IF((B{row}+C{row})=0,"N/A",D{row}/((B{row}+C{row})/2))').number_format='0.00%'
    ws.conditional_formatting.add('E17:E19',CellIsRule(operator='greaterThan',formula=['$E$5'],fill=PatternFill('solid',fgColor='FCE4D6')))
    ws['A22']='First investigation (original synthetic sample)';ws['A23']=finding
    ws['A25']='Monthly rates use exits ÷ ((opening + closing) / 2).'
    ws['A26']='Income currency is unspecified. Recruiting is an independent simulated population.'
    ws['A28']='Employee ID lookup';ws['B28']=int(source.EmployeeNumber.iloc[0]);ws['D28']='Monthly income'
    ws['E28']=f'=IFERROR(INDEX(\'Employee lookup data\'!$E$3:$E${len(source)+2},MATCH(B28,\'Employee lookup data\'!$A$3:$A${len(source)+2},0)),"Not found")'
    ws['A29']='Lookup uses original source records, including attrition-labelled employees.'
    chart=LineChart();chart.title='Simulated closing headcount';chart.y_axis.title='Employees';chart.x_axis.title='Month'
    ci=total.columns.get_loc('closing')+1
    chart.add_data(Reference(tw,min_col=ci,min_row=2,max_row=len(total)+2),titles_from_data=True)
    chart.set_categories(Reference(tw,min_col=1,min_row=3,max_row=len(total)+2));chart.width=15;chart.height=7;chart.legend=None;ws.add_chart(chart,'G7')
    ws.column_dimensions['A'].width=32
    for c in 'BCDE':ws.column_dimensions[c].width=15
    ws.column_dimensions['F'].width=3
    for c in 'GHIJKLMN':ws.column_dimensions[c].width=10
    ws.row_dimensions[23].height=42
    # Presentation-wide note, remaining cells are deliberately empty.
    ws.merge_cells('A23:N24');ws['A23'].alignment=Alignment(wrap_text=True,vertical='top')
    for w in wb:
        w.sheet_view.showGridLines=False
        for row in w:
            for cell in row:
                cell.font=Font(name='Arial',size=10,color='182B49')
                if cell.alignment.wrap_text is not True:cell.alignment=Alignment(vertical='center')
        if w!=ws:
            for c in w[2]:c.fill=PatternFill('solid',fgColor='183153');c.font=Font(name='Arial',bold=True,color='FFFFFF',size=10);c.alignment=Alignment(wrap_text=True,vertical='center')
            w.row_dimensions[2].height=32
            for col in w.columns:
                letter=col[1].column_letter;header=str(col[1].value)
                longest=max([len(header)]+[len(str(c.value)) for c in col[2:] if isinstance(c.value,str)])
                w.column_dimensions[letter].width=min(36,max(17,longest+3))
                for cell in col[2:]:
                    if isinstance(cell.value,(int,float)):cell.alignment=Alignment(horizontal='right',vertical='center',indent=1)
                    if isinstance(cell.value,float):cell.number_format='0.00%' if any(t in header for t in ['rate','share','conversion']) else '#,##0.00'
            w.print_title_rows='1:2';w.page_setup.orientation='landscape';w.page_setup.paperSize=w.PAPERSIZE_A3;w.page_setup.fitToWidth=1;w.page_setup.fitToHeight=0
        w.sheet_properties.pageSetUpPr.fitToPage=True
    for row in [16]:
        for c in ws[row][:5]:c.fill=PatternFill('solid',fgColor='183153');c.font=Font(name='Arial',bold=True,color='FFFFFF',size=10)
    ws['A2'].font=Font(name='Arial',size=16,bold=True,color='183153')
    for cell in ['B5','E5','B28']:ws[cell].fill=PatternFill('solid',fgColor='FFF2CC')
    ws.print_area='A1:N30';ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=1
    ws.sheet_properties.tabColor='183153'
    wb.calculation.fullCalcOnLoad=True
    dest=ROOT/'excel/workforce_summary.xlsx';wb.save(dest)
    # Verify exported input tables, not merely the pre-export DataFrames.
    saved=load_workbook(dest,data_only=False)
    actual=pd.DataFrame(saved['Department'].values).iloc[2:,:len(cols)].reset_index(drop=True);actual.columns=cols
    diff=max(abs(actual['closing'].astype(float)-dept.closing).max(),abs(actual['exits'].astype(float)-dept.exits).max())
    check('Excel exported counts reconcile',diff==0,f'max absolute headcount/exit difference {diff:g}')
    check('Excel requested controls',len(saved['Summary'].data_validations.dataValidation)==2 and len(saved['Summary'].conditional_formatting)>0 and 'INDEX(' in saved['Summary']['E28'].value)
    return dest
